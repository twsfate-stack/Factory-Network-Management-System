import csv
import io
import tempfile
import unittest
from pathlib import Path
from sqlalchemy import event
from sqlalchemy.exc import SQLAlchemyError
from app import create_app
from app.models import db, Switch
from app.routes.inventory_transfer import HEADERS, EXPORT_HEADERS, safe_cell


def csv_file(rows, headers=HEADERS):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=headers)
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


class InventoryTransferTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite:///" + (Path(self.temp.name)/"test.db").as_posix()})
        self.client = self.app.test_client()
        self.line = self.client.post("/api/production-lines", json={"name": "H17"}).json["data"]
        self.catalog = self.client.post("/api/switch-catalog", json={"name": "Arista 7050", "vendor": "Arista", "model": "7050"}).json["data"]

    def tearDown(self):
        with self.app.app_context():
            db.session.remove(); db.engine.dispose()
        self.temp.cleanup()

    def row(self, number=1, **extra):
        return dict(hostname=f"SW-{number}", serial_number=f"SN{number}", asset_id=f"AT{number}", switch_model="Arista 7050", production_line="H17", ip_address="172.16.10.20", status="ACTIVE") | extra

    def upload(self, raw, endpoint="preview", fingerprint=None, filename="inventory.csv"):
        data = {"file": (io.BytesIO(raw), filename)}
        if fingerprint: data["fingerprint"] = fingerprint
        return self.client.post("/api/switches/import/" + endpoint, data=data)

    def count(self):
        return self.client.get("/api/dashboard").json["total_switches"]

    def imported(self, rows):
        raw = csv_file(rows)
        review = self.upload(raw)
        self.assertEqual(review.status_code, 200, review.json)
        self.assertEqual(review.json["errors"], 0, review.json)
        response = self.upload(raw, "confirm", review.json["fingerprint"])
        self.assertEqual(response.status_code, 201, response.json)
        return response

    def test_template_and_100_rows(self):
        template = self.client.get("/api/switches/import/template")
        self.assertTrue(template.data.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(list(csv.reader(io.StringIO(template.data.decode("utf-8-sig")))), [HEADERS])
        self.assertIn("fnms_switch_inventory_template.csv", template.headers["Content-Disposition"])
        rows = [self.row(i, status=" active " if i < 80 else "offline" if i < 90 else "SPARE", notes="Thai: \u0e17\u0e14\u0e2a\u0e2d\u0e1a, \"quoted\"", firmware="4.3") for i in range(100)]
        raw = csv_file(rows); review = self.upload(raw).json
        self.assertEqual((review["total"], review["valid"], review["errors"]), (100,100,0))
        self.assertEqual(self.count(), 0)
        self.assertEqual(self.upload(raw, "confirm", review["fingerprint"]).json["imported_count"], 100)
        self.assertEqual(self.client.get("/api/dashboard").json, dict(total_switches=100, active=80, offline=10, spare=10))
        records = self.client.get("/api/switches").json
        self.assertEqual(len({r["passport_uid"] for r in records}), 100)
        self.assertTrue(all(r["created_at"].endswith("+07:00") and r["catalog_id"] == self.catalog["id"] for r in records))
        self.assertEqual(self.client.get("/api/production-lines").json[0]["total_switches"], 100)
        self.assertEqual(self.client.get("/api/switch-catalog").json[0]["switch_count"], 100)
        self.assertEqual(records[0]["firmware_version"], "4.3")

    def test_errors_duplicates_and_all_or_nothing(self):
        for field, value in [("switch_model", "missing"), ("production_line", "H99"), ("status", "REPAIR"), ("hostname", " "), ("ip_address", "bad"), ("notes", "x"*4001)]:
            raw = csv_file([self.row(1), self.row(2, **{field:value})]); preview = self.upload(raw).json
            self.assertEqual(preview["errors"], 1, preview)
            self.assertEqual(self.upload(raw, "confirm", preview["fingerprint"]).status_code, 400)
            self.assertEqual(self.count(), 0)
        for key in ("serial_number", "asset_id"):
            raw = csv_file([self.row(1, **{key:"DUP"}), self.row(2, **{key:"dup"})]); preview = self.upload(raw).json
            self.assertEqual(preview["errors"], 2)
            self.assertIn("rows 2, 3", str(preview["rows"][0]["errors"]))
        self.imported([self.row(1)])
        with self.app.app_context():
            db.session.scalar(db.select(Switch)).is_deleted = True; db.session.commit()
        for key, value in [("serial_number", "sn1"), ("asset_id", "at1")]:
            preview = self.upload(csv_file([self.row(2, **{key:value})])).json
            self.assertEqual(preview["errors"], 1)
            self.assertIn("Recycle Bin", str(preview["rows"][0]["errors"]))

    def test_stale_preview_ambiguous_match_and_tamper(self):
        raw = csv_file([self.row(1)]); preview = self.upload(raw).json
        self.assertEqual(self.upload(csv_file([self.row(2)]), "confirm", preview["fingerprint"]).status_code, 400)
        self.imported([self.row(1)])
        self.assertEqual(self.upload(raw, "confirm", preview["fingerprint"]).status_code, 400)
        self.assertEqual(self.count(), 1)
        self.client.post("/api/switch-catalog", json=dict(name="ARISTA 7050", vendor="Other", model="Other"))
        review = self.upload(csv_file([self.row(2)])).json
        self.assertIn("ambiguous", str(review["rows"][0]["errors"]))

    def test_insert_failure_rolls_back_every_row(self):
        raw = csv_file([self.row(i) for i in range(5)]); preview = self.upload(raw).json
        def fail(mapper, connection, record):
            if record.hostname == "SW-3": raise SQLAlchemyError("Simulated insertion failure")
        event.listen(Switch, "before_insert", fail)
        try:
            self.assertEqual(self.upload(raw, "confirm", preview["fingerprint"]).status_code, 503)
        finally:
            event.remove(Switch, "before_insert", fail)
        self.assertEqual(self.count(), 0)

    def test_file_validation(self):
        cases = [b"", b"hostname,status\n", b"hostname,hostname,switch_model,production_line,status\n", csv_file([]), b"hostname,switch_model,production_line,status\nSW,A,H,ACTIVE,extra\n", b'hostname,switch_model,production_line,status\n"unterminated', b"\xff\xfe", b"a\x00b", b"x"*(5*1024*1024+1)]
        for raw in cases:
            self.assertEqual(self.upload(raw).status_code, 400)
        self.assertEqual(self.upload(csv_file([self.row()]), filename="file.xlsx").status_code, 400)
        for field in ("catalog_id", "production_line_id", "created_at", "passport_uid", "image_url"):
            self.assertEqual(self.upload(csv_file([self.row() | {field:"1"}], HEADERS+[field])).status_code, 400)
        self.assertEqual(self.upload(csv_file([self.row(i) for i in range(5001)])).status_code, 400)
        self.assertEqual(self.upload(b"x"*(7*1024*1024)).status_code, 413)
        raw = b'Hostname,Switch Model,Production Line,Status\nSW, arista 7050 , h17 , active \n'
        self.assertEqual(self.upload(raw).json["valid"], 1)
        self.assertEqual(self.count(), 0)

    def test_export_all_filters_and_round_trip(self):
        rows = [self.row(i, status="ACTIVE" if i < 43 else "OFFLINE") for i in range(130)]
        self.imported(rows)
        with self.app.app_context():
            for record in db.session.scalars(db.select(Switch).order_by(Switch.id).offset(120)):
                record.is_deleted = True
            db.session.commit()
        response = self.client.get("/api/switches/export")
        self.assertTrue(response.data.startswith(b"\xef\xbb\xbf"))
        exported = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
        self.assertEqual(len(exported), 120)
        self.assertEqual(list(exported[0]), EXPORT_HEADERS)
        self.assertEqual(exported[0]["switch_model"], "Arista 7050")
        filtered = self.client.get("/api/switches/export", query_string=dict(scope="current", search="7050", production_line="H17", status="ACTIVE", vendor="Arista"))
        self.assertEqual(len(list(csv.DictReader(io.StringIO(filtered.data.decode("utf-8-sig"))))), 43)
        self.assertEqual(self.client.get("/api/switches/export?scope=current&search=no-match").status_code, 400)
        # Human columns form a new import basis; dates/Vendor/Model are intentionally omitted.
        basis = [{key: row[key] for key in HEADERS} for row in exported[:2]]
        with self.app.app_context():
            db.session.execute(db.delete(Switch)); db.session.commit()
        self.imported(basis)
        self.assertEqual(self.count(), 2)

    def test_formula_safety_and_thai(self):
        note = '\u0e17\u0e14\u0e2a\u0e2d\u0e1a, "quotes"\nsecond line'
        self.imported([self.row(1, notes=note, hostname='=SUM(1,2)')])
        result = self.client.get("/api/switches/export")
        row = next(csv.DictReader(io.StringIO(result.data.decode("utf-8-sig"))))
        self.assertEqual(row["notes"], note)
        self.assertEqual(row["hostname"], "'=SUM(1,2)")
        self.assertEqual(self.client.get("/api/switches").json[0]["hostname"], '=SUM(1,2)')
        for value in ['=1', '+1', '-1', '@X', '  =1', '\tX', '\nX']:
            self.assertTrue(safe_cell(value).startswith("'"))
        self.assertEqual(safe_cell('SW-A10'), 'SW-A10')

    def test_xlsx_all_moves_filtered_by_physical_switch_and_deleted_excluded(self):
        from openpyxl import load_workbook
        from datetime import datetime
        self.imported([self.row(1, notes="ทดสอบ", firmware="4.32"), self.row(2)])
        records = self.client.get('/api/switches').json
        target = self.client.post('/api/production-lines', json={'name': 'C20'}).json['data']
        for number in range(3):
            result = self.client.post(f"/api/switches/{records[0]['id']}/move", json={
                'production_line_id': target['id'], 'block': f'C20-0{number}',
                'hostname': '=1+1', 'notes': 'ย้ายเครื่อง'})
            self.assertEqual(result.status_code, 200)
        self.client.post(f"/api/switches/{records[1]['id']}/move", json={
            'production_line_id': target['id'], 'hostname': 'archived', 'block': 'B01'})
        with self.app.app_context():
            db.session.get(Switch, records[1]['id']).is_deleted = True
            db.session.commit()
        response = self.client.get('/api/switches/export?format=xlsx&scope=current&production_line=C20&search=SN1')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.data))
        self.assertEqual(book.sheetnames, ['Switch Inventory', 'Move History'])
        inventory, history = book.worksheets
        self.assertEqual(inventory.max_row, 2)
        self.assertEqual(history.max_row, 4)
        self.assertEqual(inventory['A2'].value, '=1+1')
        self.assertEqual(inventory['A2'].data_type, 's')
        self.assertEqual(inventory['H2'].value, 'C20')
        self.assertEqual(inventory['I2'].value, 'C20-02')
        self.assertEqual(inventory['J2'].value, '4.32')
        self.assertEqual(inventory['M2'].value, 'ทดสอบ')
        self.assertEqual(history['G2'].value, 'H17')
        self.assertEqual(history['I2'].value, 'C20')
        self.assertEqual(history['K2'].value, 'ย้ายเครื่อง')
        moves = self.client.get(f"/api/switches/{records[0]['id']}/moves").json
        expected = sorted(datetime.fromisoformat(row['moved_at']).replace(tzinfo=None, microsecond=0) for row in moves)
        actual = [row[0].replace(microsecond=0) for row in history.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(actual, expected)
        for sheet in book:
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertTrue(sheet['A1'].font.bold)
            self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        all_book = load_workbook(io.BytesIO(self.client.get('/api/switches/export?format=xlsx').data))
        self.assertEqual(all_book['Move History'].max_row, 4)
        self.assertEqual(self.client.get('/api/switches/export?format=pdf').status_code, 400)

    def test_xlsx_empty_history_and_bangkok_date(self):
        from openpyxl import load_workbook
        from datetime import datetime
        self.imported([self.row(1)])
        with self.app.app_context():
            record = db.session.scalar(db.select(Switch))
            record.created_at = datetime(2026, 9, 20, 17, 16, 32)
            db.session.commit()
        response = self.client.get('/api/switches/export?format=xlsx')
        book = load_workbook(io.BytesIO(response.data))
        self.assertEqual(book['Move History'].max_row, 1)
        self.assertEqual(book['Move History']['A1'].value, 'Move Date')
        self.assertEqual(book['Switch Inventory']['N2'].value, datetime(2026, 9, 21, 0, 16, 32))
        self.assertEqual(book['Switch Inventory']['N2'].number_format, 'yyyy-mm-dd hh:mm:ss')
