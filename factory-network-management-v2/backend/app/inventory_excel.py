"""Current inventory and persistent movement snapshots in one Excel report."""
from datetime import datetime
from io import BytesIO

from flask import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import db, SwitchMoveHistory
from .timezone import THAILAND_TZ, iso_thailand


INVENTORY_COLUMNS = [
    ('Hostname', 'hostname'), ('Serial Number', 'serial_number'),
    ('Asset ID', 'asset_id'), ('Switch Model', 'switch_model'),
    ('Vendor', 'vendor'), ('Model', 'model'), ('IP Address', 'ip_address'),
    ('Production Line', 'production_line'), ('Block', 'block'),
    ('Firmware', 'firmware'), ('MAC Address', 'mac_address'),
    ('Status', 'status'), ('Notes', 'notes'),
    ('Created Date', 'created_at'), ('Updated Date', 'updated_at'),
]
MOVE_HEADERS = ['Move Date', 'Serial Number', 'Asset ID', 'Switch Model',
                'Old Hostname', 'New Hostname', 'From Production Line',
                'From Block', 'To Production Line', 'To Block', 'Note']


def excel_date(value):
    # Excel dates have no timezone: convert once, then store Bangkok wall time.
    return datetime.fromisoformat(value).astimezone(THAILAND_TZ).replace(tzinfo=None) if value else None


def add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for values in rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                # Store untrusted data as literal text, never as Excel formulas.
                cell.data_type = 's'
            if isinstance(cell.value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='393E46')
        cell.fill = PatternFill('solid', fgColor='EEEEEE')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 42 if header in ('Notes', 'Note') else 24
    return sheet


def text_value(value):
    # XML 1.0 cannot represent these controls; preserve Thai and normal whitespace.
    return ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value


def excel_download(records, filename):
    workbook = Workbook()
    workbook.remove(workbook.active)
    inventory = []
    for record in records:
        inventory.append([excel_date(record.get(key)) if key in ('created_at', 'updated_at')
                          else text_value(record.get(key)) for _, key in INVENTORY_COLUMNS])
    add_sheet(workbook, 'Switch Inventory', [label for label, _ in INVENTORY_COLUMNS], inventory)
    by_id = {record['id']: record for record in records}
    moves = []
    # Chunk IDs to stay below SQLite parameter limits; retain every saved move.
    ids = list(by_id)
    for start in range(0, len(ids), 500):
        moves.extend(db.session.scalars(db.select(SwitchMoveHistory).where(
            SwitchMoveHistory.switch_id.in_(ids[start:start + 500]))))
    moves.sort(key=lambda move: (move.moved_at, move.id))
    history = []
    for move in moves:
        record = by_id[move.switch_id]
        history.append([excel_date(iso_thailand(move.moved_at)),
                        record.get('serial_number'), record.get('asset_id'), record['switch_model'],
                        move.from_hostname, move.to_hostname, move.from_production_line_name,
                        move.from_block, move.to_production_line_name, move.to_block, move.notes])
    add_sheet(workbook, 'Move History', MOVE_HEADERS,
              ([text_value(value) for value in row] for row in history))
    output = BytesIO()
    workbook.save(output)
    response = Response(output.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.headers['Cache-Control'] = 'no-store'
    return response
